
from __future__ import annotations

import io
import re
from pathlib import Path

import fitz
from docx import Document
from pptx import Presentation


MAX_UPLOAD_BYTES = (
    10 * 1024 * 1024
)

RESUME_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
}

LEARNING_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".srt",
    ".vtt",
}


class ExtractionError(ValueError):
    """Raised when an uploaded file cannot be safely parsed."""
    pass


def _clean_text(
    text: str,
) -> str:

    text = text.replace(
        "\x00",
        " ",
    )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def _extract_pdf(
    content: bytes,
) -> str:

    try:

        doc = fitz.open(
            stream=content,
            filetype="pdf",
        )

    except Exception as exc:

        raise ExtractionError(
            "The PDF could not be opened. "
            "Please upload a valid PDF file."
        ) from exc

    parts: list[str] = []

    try:

        for page in doc:

            parts.append(
                page.get_text("text")
            )

    finally:

        doc.close()

    return "\n".join(
        parts
    )


def _extract_docx(
    content: bytes,
) -> str:

    try:

        doc = Document(
            io.BytesIO(content)
        )

    except Exception as exc:

        raise ExtractionError(
            "The DOCX file could not be opened."
        ) from exc

    lines: list[str] = []

    for paragraph in doc.paragraphs:

        if paragraph.text.strip():

            lines.append(
                paragraph.text
            )

    for table in doc.tables:

        for row in table.rows:

            values = [

                cell.text.strip()

                for cell in row.cells

                if cell.text.strip()

            ]

            if values:

                lines.append(
                    " | ".join(values)
                )

    return "\n".join(
        lines
    )


def _extract_pptx(
    content: bytes,
) -> str:

    try:

        prs = Presentation(
            io.BytesIO(content)
        )

    except Exception as exc:

        raise ExtractionError(
            "The PPTX file could not be opened."
        ) from exc

    parts: list[str] = []

    for (
        slide_index,
        slide,
    ) in enumerate(
        prs.slides,
        start=1,
    ):

        slide_text: list[str] = []

        for shape in slide.shapes:

            if (
                hasattr(shape, "text")
                and shape.text.strip()
            ):

                slide_text.append(
                    shape.text.strip()
                )

        if slide_text:

            parts.append(

                f"[Slide {slide_index}]\n"

                + "\n".join(
                    slide_text
                )

            )

    return "\n\n".join(
        parts
    )


def extract_text(
    filename: str,
    content: bytes,
    *,
    purpose: str,
) -> str:
    """
    Extract text from a supported upload.

    purpose:
        resume
        learning
    """

    if not filename:

        raise ExtractionError(
            "Uploaded file must have a filename."
        )

    if not content:

        raise ExtractionError(
            "The uploaded file is empty."
        )

    if len(content) > MAX_UPLOAD_BYTES:

        raise ExtractionError(
            "File is larger than the "
            "10 MB prototype upload limit."
        )

    extension = (
        Path(filename)
        .suffix
        .lower()
    )

    allowed = (

        RESUME_EXTENSIONS

        if purpose == "resume"

        else LEARNING_EXTENSIONS

    )

    if extension not in allowed:

        supported = ", ".join(
            sorted(allowed)
        )

        raise ExtractionError(

            f"Unsupported file type "
            f"'{extension}'. "
            f"Supported: {supported}"

        )

    if extension == ".pdf":

        text = _extract_pdf(
            content
        )

    elif extension == ".docx":

        text = _extract_docx(
            content
        )

    elif extension == ".pptx":

        text = _extract_pptx(
            content
        )

    else:

        try:

            text = content.decode(
                "utf-8-sig"
            )

        except UnicodeDecodeError:

            text = content.decode(
                "latin-1",
                errors="replace",
            )

    text = _clean_text(
        text
    )

    if len(text) < 40:

        if extension == ".pdf":

            raise ExtractionError(
                "Very little selectable text was found "
                "in this PDF. It may be a scanned/image-only "
                "PDF. For the SIH prototype, use a digital/text "
                "PDF or DOCX. OCR can be added in the "
                "production phase."
            )

        raise ExtractionError(
            "The document does not contain "
            "enough readable text to analyze."
        )

    return text